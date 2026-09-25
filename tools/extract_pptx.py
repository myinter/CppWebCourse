#!/usr/bin/env python3
"""
把 C++ 教程的 PPT 导出成结构化 JSON，供人工校订与撰写网页内容使用。

导出内容：
  - 每页的标题、正文文本块（按位置排序，还原阅读顺序与分组）
  - 演讲者备注（逐字）
  - 识别出的代码块（含原文与缺陷标记）

用法：python3 tools/extract_pptx.py
输出：tools/out/preface.json, ch1.json … ch7.json  +  tools/out/summary.txt
"""
import json
import os
import re
import glob
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.dirname(os.path.dirname(HERE))  # C++/
OUT = os.path.join(HERE, "out")

# 按文件名匹配，统一编号
CHAPTERS = [
    ("前言", 0, "课程介绍", "为什么要学 C++、这门课怎么上"),
    ("第1章", 1, "数据类型与变量常量", "基本类型、变量、常量、输入输出"),
    ("第2章", 2, "运算符与表达式", "算术、关系、逻辑、位运算与优先级"),
    ("第3章", 3, "控制流", "分支、循环、跳转语句"),
    ("第4章", 4, "函数", "定义、参数、返回值、递归、重载"),
    ("第5章", 5, "数组与字符串", "一维/二维数组、字符数组与 string"),
    ("第6章", 6, "指针与函数", "指针、引用、指针与函数、动态内存"),
    ("第7章", 7, "结构体与共用体以及枚举", "struct / union / enum / typedef"),
]

# C++ 代码特征
CODE_SIGNALS = re.compile(
    r"(#include)|(#define)|(\bstd::)|(\busing namespace)|(\bcout\b)|(\bcin\b)|"
    r"(\bprintf\b)|(\bscanf\b)|(\bint\s+main)|(\bint\s+\w+\s*[=;(])|"
    r"(\bfor\s*\()|(\bwhile\s*\()|(\bif\s*\()|(\belse\b)|(\bswitch\s*\()|"
    r"(\bstruct\s+\w)|(\bunion\s+\w)|(\benum\s+\w)|(\bclass\s+\w)|"
    r"(\bvoid\s+\w+\s*\()|(\breturn\b)|(\bconst\b)|(\bdouble\b)|(\bfloat\b)|"
    r"(\bchar\b)|(\bbool\b)|(\bstring\b)|(\bnew\b)|(\bdelete\b)|(\bnullptr\b)"
)
CXX_STRONG = re.compile(r"(#include)|(;)|(std::)|(cout|cin)|(\bint\s+main)|(\breturn\b)")

# PPT 导出代码时常见的缺陷
DEFECTS = {
    "fullwidth_punct": re.compile(r"[，。！？；：（）【】、“”‘’]"),
    "smart_quote": re.compile(r"[“”‘’]"),
    "missing_semicolon": None,   # 需要语义判断，不自动标记
    "ellipsis": re.compile(r"\.\.\s*\.\.|…"),
    "tight_arrow": re.compile(r"\w->\w"),
}


def para_text(p):
    """还原段落文本，保留行内换行符。"""
    parts = []
    for child in p._p:
        tag = child.tag.split("}")[-1]
        if tag == "br":
            parts.append("\n")
        elif tag == "r":
            t = child.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}t")
            if t is not None and t.text:
                parts.append(t.text)
    return "".join(parts)


def shape_paragraphs(sh):
    if not sh.has_text_frame:
        return []
    out = []
    for p in sh.text_frame.paragraphs:
        txt = para_text(p)
        if txt.strip():
            out.append({"text": txt, "level": p.level or 0})
    return out


def walk(shapes, depth=0):
    for sh in shapes:
        yield sh, depth
        if sh.shape_type == 6:  # GROUP
            try:
                yield from walk(sh.shapes, depth + 1)
            except Exception:
                pass


def looks_like_code(text):
    """判断一段文本是否为代码。"""
    strong = len(CXX_STRONG.findall(text))
    weak = len(CODE_SIGNALS.findall(text))
    # 有分号或 #include 且命中多个特征
    return (strong >= 2 and weak >= 3) or (weak >= 6)


def mark_defects(code):
    found = []
    for name, pat in DEFECTS.items():
        if pat and pat.search(code):
            found.append(name)
    return found


def extract(path, meta):
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        blocks = []
        for sh, depth in walk(slide.shapes):
            paras = shape_paragraphs(sh)
            if not paras:
                continue
            try:
                pos = {
                    "x": round(Emu(sh.left).inches, 2) if sh.left is not None else None,
                    "y": round(Emu(sh.top).inches, 2) if sh.top is not None else None,
                    "w": round(Emu(sh.width).inches, 2) if sh.width is not None else None,
                    "h": round(Emu(sh.height).inches, 2) if sh.height is not None else None,
                }
            except Exception:
                pos = {"x": None, "y": None, "w": None, "h": None}

            text = "\n".join(p["text"] for p in paras)
            is_code = looks_like_code(text)
            entry = {
                "text": text,
                "lines": [p["text"] for p in paras],
                "pos": pos,
                "depth": depth,
                "is_code": is_code,
            }
            if is_code:
                entry["defects"] = mark_defects(text)
            blocks.append(entry)

        blocks.sort(key=lambda b: (b["pos"]["y"] if b["pos"]["y"] is not None else 999,
                                   b["pos"]["x"] if b["pos"]["x"] is not None else 999))

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        title = None
        for b in blocks:
            if not b["is_code"] and b["pos"]["y"] is not None and b["pos"]["y"] < 2.2:
                title = b["lines"][0]
                break
        if title is None and blocks:
            title = blocks[0]["lines"][0]

        slides.append({
            "n": idx,
            "title_guess": title,
            "notes": notes,
            "blocks": blocks,
            "code_blocks": [b for b in blocks if b["is_code"]],
        })

    return {
        "chapter": meta["num"], "name": meta["name"], "subtitle": meta["subtitle"],
        "source": os.path.basename(path), "slide_count": len(slides), "slides": slides
    }


def main():
    os.makedirs(OUT, exist_ok=True)
    files = sorted(glob.glob(os.path.join(SRC, "*.pptx")))

    ordered = []
    for prefix, num, name, subtitle in CHAPTERS:
        key = prefix.replace("第", "").replace("章", "")
        for f in files:
            base = os.path.basename(f)
            if prefix == "前言" and "前言" in base:
                ordered.append((f, {"num": num, "name": name, "subtitle": subtitle}))
                break
            if f"第{key}章" in base:
                ordered.append((f, {"num": num, "name": name, "subtitle": subtitle}))
                break

    summary = []
    for path, meta in ordered:
        data = extract(path, meta)
        tag = "preface" if meta["num"] == 0 else f"ch{meta['num']}"
        dest = os.path.join(OUT, f"{tag}.json")
        with open(dest, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        ncode = sum(len(s["code_blocks"]) for s in data["slides"])
        nnotes = sum(1 for s in data["slides"] if s["notes"])
        summary.append(
            f"{tag:<8} {meta['name'][:14]:<16} {data['slide_count']:>3} 页  "
            f"代码块 {ncode:>2}  备注 {nnotes:>2}/{data['slide_count']}")
        # 报告缺陷
        defects = {}
        for s in data["slides"]:
            for b in s["code_blocks"]:
                for d in b.get("defects", []):
                    defects[d] = defects.get(d, 0) + 1
        if defects:
            summary.append("          缺陷: " + ", ".join(f"{k}×{v}" for k, v in defects.items()))

    out = "\n".join(summary)
    with open(os.path.join(OUT, "summary.txt"), "w", encoding="utf-8") as f:
        f.write(out + "\n")
    print(out)
    print(f"\n输出目录: {OUT}")


if __name__ == "__main__":
    main()
